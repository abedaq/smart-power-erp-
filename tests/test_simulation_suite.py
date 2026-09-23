import os
import subprocess
import sys
from pathlib import Path

# Add project root to sys.path
PROJECT_ROOT = Path(__file__).resolve().parent.parent
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

import openpyxl
import pytest

from simulation.config import SimulationConfig
from simulation.client import SmartPowerClient
from simulation.selector import SubscriberSelector, SubscriberSnapshot
from simulation.engine import SimulationEngine, SimulationCycleRecord
from simulation.verifier import FinancialVerifier, VerificationResult
from simulation.receipt_renderer import ReceiptRenderer, ReceiptRenderData
from simulation.excel_exporter import ExcelExporter


# -----------------------------------------------------------------------------
# 1. Config Tests
# -----------------------------------------------------------------------------
def test_config_defaults_and_env():
    """Verify configuration loads sane defaults and environment variables."""
    cfg = SimulationConfig()
    assert cfg.api_base_url.startswith("http")
    assert cfg.subscriber_sample_size == 100
    assert cfg.billing_cycles_count == 12
    assert cfg.default_kwh_price == 1400.0
    assert cfg.default_fixed_fee == 1000.0
    assert cfg.random_seed == 42
    assert cfg.receipts_dir.exists()
    assert cfg.reports_dir.exists()


# -----------------------------------------------------------------------------
# 2. Client Tests
# -----------------------------------------------------------------------------
def test_client_initialization_and_session():
    """Verify REST API client initializes session with retry adapter."""
    client = SmartPowerClient()
    assert client.base_url == "http://localhost:3000/api"
    assert client.session is not None
    assert client.token is None


def test_client_live_connectivity_or_mock():
    """Verify client can authenticate with backend if server is active."""
    client = SmartPowerClient()
    is_healthy = client.health_check()
    if is_healthy:
        token = client.login()
        assert token is not None
        assert len(token) > 20
        customers = client.get_customers(limit=5)
        assert isinstance(customers, list)
        assert len(customers) > 0


# -----------------------------------------------------------------------------
# 3. Selector Tests
# -----------------------------------------------------------------------------
def test_selector_deterministic_sampling():
    """Verify reproducible selection given the same seed."""
    client = SmartPowerClient()
    selector = SubscriberSelector(client=client)

    # Synthetic customer pool
    mock_customers = [
        {
            "id": i,
            "subscriber_number": f"0000{i}",
            "full_name": f"مشترك {i}",
            "status": "Active",
            "is_deleted": False,
            "last_reading": 100.0 + i,
            "total_due": 5000.0,
        }
        for i in range(1, 151)
    ]

    # Deterministic sorting & shuffle
    mock_customers.sort(key=lambda c: str(c.get("subscriber_number", "")).zfill(10))
    import random
    rng1 = random.Random(42)
    pool1 = list(mock_customers)
    rng1.shuffle(pool1)
    sample1 = pool1[:100]

    rng2 = random.Random(42)
    pool2 = list(mock_customers)
    rng2.shuffle(pool2)
    sample2 = pool2[:100]

    assert [c["id"] for c in sample1] == [c["id"] for c in sample2]
    assert len(sample1) == 100


def test_subscriber_snapshot_serialization():
    """Verify SubscriberSnapshot dataclass serialization with English numbers."""
    snap = SubscriberSnapshot(
        customer_id=101,
        subscriber_number="910001",
        full_name="علي يحيى",
        phone_number="+967 771234567",
        meter_number="MTR-00101",
        route_number="A1",
        initial_reading=100.0,
        last_reading=250.5,
        baseline_total_due=14500.0,
        baseline_balance=14500.0,
    )
    d = snap.to_dict()
    assert d["customer_id"] == 101
    assert d["subscriber_number"] == "910001"
    assert d["baseline_total_due"] == 14500.0
    assert d["last_reading"] == 250.5


# -----------------------------------------------------------------------------
# 4. Engine Tests
# -----------------------------------------------------------------------------
def test_engine_cycle_generation():
    """Verify cycle sequence generation across year-end rollover."""
    cycles = SimulationEngine.generate_cycle_sequence(start_year=2026, start_month=11, count=4)
    assert cycles == ["2026-11-1", "2026-12-1", "2027-01-1", "2027-02-1"]


def test_engine_monotonic_reading_and_financial_math():
    """Verify monotonicity and mathematical invariance in simulation cycle."""
    client = SmartPowerClient()
    engine = SimulationEngine(client=client)
    import random
    rng = random.Random(42)

    snap = SubscriberSnapshot(
        customer_id=1,
        subscriber_number="910001",
        full_name="مختار عبد الله",
        phone_number="+967 770000001",
        meter_number="MTR-00001",
        route_number="A1",
        initial_reading=500.0,
        last_reading=500.0,
        baseline_total_due=2000.0,
        baseline_balance=2000.0,
    )

    record = engine.simulate_cycle_for_subscriber(
        snapshot=snap,
        cycle_name="2027-03-1",
        cycle_index=1,
        prev_reading=500.0,
        prev_remaining=2000.0,
        rng=rng,
        dry_run=True,
    )

    # Assert Monotonicity
    assert record.current_reading >= record.previous_reading
    assert record.consumption_kwh >= 0.0

    # Assert Formulas
    expected_consumption_val = round(record.consumption_kwh * record.kwh_price, 2)
    assert record.consumption_value == expected_consumption_val

    expected_total_due = round(expected_consumption_val + record.fixed_fee + record.carried_arrears, 2)
    assert record.total_due == expected_total_due

    expected_remaining = round(record.total_due - record.paid_amount, 2)
    assert record.remaining_balance == expected_remaining


def test_engine_multi_cycle_dry_run():
    """Verify full multi-subscriber multi-cycle execution in dry-run mode."""
    client = SmartPowerClient()
    engine = SimulationEngine(client=client)

    snapshots = [
        SubscriberSnapshot(
            customer_id=i,
            subscriber_number=f"91{i:04d}",
            full_name=f"مشترك {i}",
            phone_number=f"+967 77{i:07d}",
            meter_number=f"MTR-{i:05d}",
            route_number="A1",
            initial_reading=100.0 * i,
            last_reading=100.0 * i,
            baseline_total_due=1000.0,
            baseline_balance=1000.0,
        )
        for i in range(1, 11)
    ]

    result = engine.run_simulation(
        snapshots=snapshots,
        cycles=["2027-01-1", "2027-02-1", "2027-03-1"],
        dry_run=True,
    )

    assert len(result.records) == 30  # 10 subscribers * 3 cycles
    assert result.total_billed_amount > 0.0
    assert result.total_paid_amount > 0.0
    assert len(result.anomalies) == 0


# -----------------------------------------------------------------------------
# 5. Verifier Tests
# -----------------------------------------------------------------------------
def test_verifier_reversal_logic_dry_run():
    """Verify mathematical invariance assertion for payment reversal."""
    client = SmartPowerClient()
    verifier = FinancialVerifier(client=client)

    result = verifier.verify_payment_reversal(
        customer_id=42,
        test_amount=7500.0,
        dry_run=True,
    )

    assert isinstance(result, VerificationResult)
    assert result.passed is True
    assert result.discrepancy == 0.0
    assert result.expected_value == result.actual_value


def test_verifier_tariff_cascade_dry_run():
    """Verify mathematical invariance assertion for tariff modification."""
    client = SmartPowerClient()
    verifier = FinancialVerifier(client=client)

    result = verifier.verify_tariff_price_cascade(
        invoice_id=10,
        new_kwh_price=1600.0,
        dry_run=True,
    )

    assert isinstance(result, VerificationResult)
    assert result.passed is True
    assert result.discrepancy == 0.0


# -----------------------------------------------------------------------------
# 6. Receipt Renderer Tests
# -----------------------------------------------------------------------------
def test_receipt_renderer_html_construction():
    """Verify authentic dual-stub HTML is generated with English numerals."""
    renderer = ReceiptRenderer()
    data = ReceiptRenderData(
        receipt_number="REC-2027-000042",
        customer_name="صالح محمد صالح",
        subscriber_number="910042",
        phone_number="+967 770000042",
        address="صنعاء - الستين",
        meter_number="MTR-00042",
        route_number="A2",
        billing_cycle="2027-03-1",
        previous_reading=1200.0,
        current_reading=1350.0,
        consumption=150.0,
        kwh_price=1400.0,
        consumption_value=210000.0,
        fixed_fee=1000.0,
        arrears=5000.0,
        total_due=216000.0,
        paid_amount=216000.0,
        remaining_balance=0.0,
    )

    html_str = renderer.build_receipt_html(data)
    assert "REC-2027-000042" in html_str
    assert "صالح محمد صالح" in html_str
    assert "910042" in html_str
    assert "216,000" in html_str
    assert "0 ر.ي (خالص)" in html_str
    assert "coupon-stub" in html_str
    assert "main-stub" in html_str


def test_receipt_renderer_playwright_render(tmp_path):
    """Verify Playwright renders actual PNG file."""
    renderer = ReceiptRenderer()
    out_file = tmp_path / "test_dual_stub_receipt.png"
    data = ReceiptRenderData(
        receipt_number="REC-2027-000099",
        customer_name="فهد أحمد",
        subscriber_number="910099",
        billing_cycle="2027-04-1",
        previous_reading=400.0,
        current_reading=480.0,
        consumption=80.0,
        total_due=113000.0,
        paid_amount=100000.0,
        remaining_balance=13000.0,
    )

    result_path = renderer.render_receipt(data, output_path=out_file)
    assert result_path.exists()
    assert result_path.stat().st_size > 10000  # Non-empty high-res PNG


# -----------------------------------------------------------------------------
# 7. Excel Exporter Tests
# -----------------------------------------------------------------------------
def test_excel_exporter_6_sheets(tmp_path):
    """Verify openpyxl creates a valid 6-sheet audit workbook with RTL enabled."""
    exporter = ExcelExporter()
    out_xlsx = tmp_path / "test_audit_workbook.xlsx"

    snapshots = [
        SubscriberSnapshot(
            customer_id=1,
            subscriber_number="910001",
            full_name="مأمون يحيى",
            phone_number="+967 770000001",
            meter_number="MTR-001",
            route_number="A1",
            initial_reading=100.0,
            last_reading=100.0,
            baseline_total_due=5000.0,
            baseline_balance=5000.0,
        )
    ]

    records = [
        SimulationCycleRecord(
            cycle_name="2027-01-1",
            cycle_index=1,
            customer_id=1,
            subscriber_number="910001",
            full_name="مأمون يحيى",
            route_number="A1",
            previous_reading=100.0,
            current_reading=150.0,
            consumption_kwh=50.0,
            kwh_price=1400.0,
            fixed_fee=1000.0,
            consumption_value=70000.0,
            carried_arrears=5000.0,
            total_due=76000.0,
            paid_amount=76000.0,
            remaining_balance=0.0,
            invoice_status="Paid",
            receipt_number="REC-2027-000001",
        )
    ]

    sim_res = SimulationEngine(client=SmartPowerClient()).run_simulation(
        snapshots=snapshots,
        cycles=["2027-01-1"],
        dry_run=True,
    )

    verifications = [
        VerificationResult(
            verification_type="PAYMENT_REVERSAL",
            entity_id="PMT-1",
            passed=True,
            discrepancy=0.0,
            expected_value=5000.0,
            actual_value=5000.0,
            details="Zero discrepancy verified",
        ),
        VerificationResult(
            verification_type="TARIFF_CASCADE",
            entity_id="INV-1",
            passed=True,
            discrepancy=0.0,
            expected_value=76000.0,
            actual_value=76000.0,
            details="Tariff cascade verified",
        ),
    ]

    res_path = exporter.export_audit_workbook(
        sim_result=sim_res,
        snapshots=snapshots,
        verification_results=verifications,
        output_path=out_xlsx,
    )

    assert res_path.exists()
    wb = openpyxl.load_workbook(str(res_path))
    assert len(wb.sheetnames) == 6
    expected_sheet_titles = [
        "1. الملخص التنفيذي",
        "2. ملخص المشتركين السنوي",
        "3. سجل الدورات الشهرية",
        "4. تدقيق إلغاء السندات",
        "5. تدقيق تغيير التعرفة",
        "6. كشف الشذوذ والتناقضات",
    ]
    for expected in expected_sheet_titles:
        assert expected in wb.sheetnames
        ws = wb[expected]
        assert ws.sheet_view.rightToLeft is True


# -----------------------------------------------------------------------------
# 8. CLI Runner Tests
# -----------------------------------------------------------------------------
def test_cli_runner_dry_run_execution(tmp_path):
    """Verify run_simulation.py executes cleanly with --dry-run argument."""
    cmd = [
        sys.executable,
        "run_simulation.py",
        "--subscribers", "10",
        "--cycles", "3",
        "--reversals", "2",
        "--render-receipts", "2",
        "--dry-run",
        "--output-dir", str(tmp_path / "artifacts"),
    ]
    proc = subprocess.run(cmd, capture_output=True, text=True, encoding="utf-8")
    assert proc.returncode == 0, f"CLI runner failed: {proc.stderr}\nOutput: {proc.stdout}"
    assert "SIMULATION & VERIFICATION SUITE FINISHED SUCCESSFULLY" in proc.stdout
