"""
Multi-sheet financial audit Excel workbook generator for SmartPower ERP.
Generates a comprehensive 6-sheet audit workbook using openpyxl with strict English numerals and RTL layout.
"""

import logging
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Dict, List, Optional
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from simulation.config import SimulationConfig, config as default_config
from simulation.engine import SimulationCycleRecord, SimulationRunResult
from simulation.selector import SubscriberSnapshot
from simulation.verifier import VerificationResult

logger = logging.getLogger("simulation.excel_exporter")


class ExcelExporter:
    """Exports multi-sheet financial audit spreadsheets with openpyxl."""

    def __init__(self, config: Optional[SimulationConfig] = None):
        self.config = config or default_config

        # Color palette
        self.header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
        self.header_font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
        self.kpi_title_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
        self.kpi_title_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
        self.kpi_val_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
        self.kpi_val_font = Font(name="Segoe UI", size=14, bold=True, color="1E3A8A")

        self.zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
        self.white_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
        self.green_fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
        self.green_font = Font(name="Segoe UI", size=10, bold=True, color="047857")
        self.red_fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
        self.red_font = Font(name="Segoe UI", size=10, bold=True, color="B91C1C")

        thin_side = Side(style="thin", color="CBD5E1")
        self.cell_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
        double_side = Side(style="double", color="0F172A")
        self.total_border = Border(top=thin_side, bottom=double_side, left=thin_side, right=thin_side)

        self.align_center = Alignment(horizontal="center", vertical="center", wrap_text=False)
        self.align_right = Alignment(horizontal="right", vertical="center", wrap_text=False)
        self.align_left = Alignment(horizontal="left", vertical="center", wrap_text=False)

    def _apply_sheet_view_rtl(self, ws: openpyxl.worksheet.worksheet.Worksheet):
        """Set right-to-left layout and default grid visibility."""
        ws.sheet_view.rightToLeft = True
        ws.views.sheetView[0].showGridLines = True

    def _auto_adjust_columns(self, ws: openpyxl.worksheet.worksheet.Worksheet, min_width: int = 12):
        """Auto-adjust column widths based on maximum string lengths."""
        for col in ws.columns:
            col_letter = get_column_letter(col[0].column)
            max_len = 0
            for cell in col:
                val = str(cell.value or "")
                if val:
                    # Treat multiline or long text gracefully
                    lines = val.split("\n")
                    line_len = max(len(l) for l in lines)
                    if line_len > max_len:
                        max_len = line_len
            ws.column_dimensions[col_letter].width = max(max_len + 4, min_width)

    # -------------------------------------------------------------------------
    # Sheet 1: Executive Summary & Audit Verdict
    # -------------------------------------------------------------------------
    def _create_executive_summary_sheet(
        self,
        wb: openpyxl.Workbook,
        sim_result: SimulationRunResult,
        verification_results: List[VerificationResult],
    ):
        ws = wb.create_sheet(title="1. الملخص التنفيذي")
        self._apply_sheet_view_rtl(ws)

        # Title Banner
        ws.merge_cells("A1:G1")
        top_cell = ws["A1"]
        top_cell.value = "تقرير التدقيق المالي ومحاكاة عمليات الفوترة والتحصيل (SmartPower ERP)"
        top_cell.font = Font(name="Segoe UI", size=14, bold=True, color="FFFFFF")
        top_cell.fill = self.kpi_title_fill
        top_cell.alignment = self.align_center
        ws.row_dimensions[1].height = 36

        # Subtitle Info
        ws.merge_cells("A2:G2")
        sub_cell = ws["A2"]
        sub_cell.value = f"تاريخ التشغيل: {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M:%S UTC')} | بيئة الاختبار: Standalone Go Backend + PostgreSQL 18"
        sub_cell.font = Font(name="Segoe UI", size=9, italic=True, color="64748B")
        sub_cell.alignment = self.align_center
        ws.row_dimensions[2].height = 20

        # KPI Cards (2 rows of 3 metrics)
        kpis = [
            ("عدد المشتركين المختبرين", f"{sim_result.subscribers_count} مشترك"),
            ("عدد الدورات المحاكاة", f"{sim_result.cycles_count} دورة"),
            ("إجمالي حركات الفوترة", f"{len(sim_result.records):,} حركة"),
            ("إجمالي المبالغ المفوترة", f"{sim_result.total_billed_amount:,.2f} YER"),
            ("إجمالي المبالغ المحصلة", f"{sim_result.total_paid_amount:,.2f} YER"),
            ("إجمالي المتأخرات المتبقية", f"{sim_result.total_remaining_arrears:,.2f} YER"),
        ]

        row_start = 4
        for idx, (label, val) in enumerate(kpis):
            r = row_start + (idx // 3) * 3
            c_start = 1 + (idx % 3) * 2
            c_end = c_start + 1

            ws.merge_cells(start_row=r, start_column=c_start, end_row=r, end_column=c_end)
            lbl_cell = ws.cell(row=r, column=c_start, value=label)
            lbl_cell.fill = self.kpi_title_fill
            lbl_cell.font = Font(name="Segoe UI", size=9, bold=True, color="FFFFFF")
            lbl_cell.alignment = self.align_center

            ws.merge_cells(start_row=r + 1, start_column=c_start, end_row=r + 1, end_column=c_end)
            val_cell = ws.cell(row=r + 1, column=c_start, value=val)
            val_cell.fill = self.kpi_val_fill
            val_cell.font = self.kpi_val_font
            val_cell.alignment = self.align_center

        # Final Verdict Box
        verdict_row = 11
        ws.merge_cells(f"A{verdict_row}:G{verdict_row}")
        v_title = ws[f"A{verdict_row}"]
        v_title.value = "القرار المحاسبي والرقابي النهائي (Final Audit Verdict)"
        v_title.fill = self.header_fill
        v_title.font = self.header_font
        v_title.alignment = self.align_center

        reversals_passed = all(r.passed for r in verification_results) if verification_results else True
        zero_anomalies = len(sim_result.anomalies) == 0
        all_passed = reversals_passed and zero_anomalies

        ws.merge_cells(f"A{verdict_row+1}:G{verdict_row+2}")
        v_cell = ws[f"A{verdict_row+1}"]
        if all_passed:
            v_cell.value = "نظام الفوترة والحسابات معتمد بنجاح 100% — خالٍ من أي انحرافات مالية (AUDIT VERDICT: PASSED)"
            v_cell.fill = self.green_fill
            v_cell.font = Font(name="Segoe UI", size=12, bold=True, color="047857")
        else:
            v_cell.value = "تم رصد انحرافات حسابية أثناء المحاكاة (AUDIT VERDICT: FAILED - REVISION REQUIRED)"
            v_cell.fill = self.red_fill
            v_cell.font = Font(name="Segoe UI", size=12, bold=True, color="B91C1C")
        v_cell.alignment = self.align_center

        self._auto_adjust_columns(ws, min_width=16)

    # -------------------------------------------------------------------------
    # Sheet 2: Subscriber Annual Summary (100 Subscribers)
    # -------------------------------------------------------------------------
    def _create_subscriber_summary_sheet(
        self,
        wb: openpyxl.Workbook,
        sim_result: SimulationRunResult,
        snapshots: List[SubscriberSnapshot],
    ):
        ws = wb.create_sheet(title="2. ملخص المشتركين السنوي")
        self._apply_sheet_view_rtl(ws)

        headers = [
            "م", "رقم المشترك", "اسم المشترك", "خط السير",
            "الرصيد الافتتاحي (ر.ي)", "إجمالي الاستهلاك (ك.و)", "إجمالي المفوتر (ر.ي)",
            "إجمالي المسدد (ر.ي)", "الرصيد الختامي المحسوب (ر.ي)",
            "فارق المطابقة (ر.ي)", "حالة المطابقة"
        ]

        ws.row_dimensions[1].height = 28
        for col_idx, h in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col_idx, value=h)
            cell.fill = self.header_fill
            cell.font = self.header_font
            cell.alignment = self.align_center
            cell.border = self.cell_border

        # Aggregate records by subscriber
        sub_records: Dict[int, List[SimulationCycleRecord]] = {}
        for r in sim_result.records:
            sub_records.setdefault(r.customer_id, []).append(r)

        row_num = 2
        for idx, snap in enumerate(snapshots, start=1):
            cid = snap.customer_id
            records = sub_records.get(cid, [])

            total_cons = sum(r.consumption_kwh for r in records)
            total_billed = sum(r.total_due for r in records)
            total_paid = sum(r.paid_amount for r in records)
            closing_balance = records[-1].remaining_balance if records else snap.baseline_total_due

            expected_closing = round(snap.baseline_total_due + sum(r.consumption_value + r.fixed_fee - r.paid_amount for r in records), 2)
            variance = round(abs(closing_balance - expected_closing), 2)
            match_status = "MATCH" if variance < 0.05 else "MISMATCH"

            ws.row_dimensions[row_num].height = 20
            fill = self.zebra_fill if row_num % 2 == 0 else self.white_fill

            vals = [
                idx,
                snap.subscriber_number,
                snap.full_name,
                snap.route_number,
                round(snap.baseline_total_due, 2),
                round(total_cons, 2),
                round(total_billed, 2),
                round(total_paid, 2),
                round(closing_balance, 2),
                variance,
                match_status,
            ]

            for c_idx, v in enumerate(vals, start=1):
                cell = ws.cell(row=row_num, column=c_idx, value=v)
                cell.fill = fill
                cell.border = self.cell_border
                cell.font = Font(name="Segoe UI", size=9)

                if c_idx in (1, 2, 4, 11):
                    cell.alignment = self.align_center
                elif c_idx == 3:
                    cell.alignment = self.align_right
                else:
                    cell.alignment = self.align_left
                    cell.number_format = "#,##0.00"

                if c_idx == 11:
                    cell.fill = self.green_fill if match_status == "MATCH" else self.red_fill
                    cell.font = self.green_font if match_status == "MATCH" else self.red_font

            row_num += 1

        self._auto_adjust_columns(ws, min_width=14)

    # -------------------------------------------------------------------------
    # Sheet 3: Monthly Cycle Ledger (1,200 Rows)
    # -------------------------------------------------------------------------
    def _create_cycle_ledger_sheet(
        self,
        wb: openpyxl.Workbook,
        sim_result: SimulationRunResult,
    ):
        ws = wb.create_sheet(title="3. سجل الدورات الشهرية")
        self._apply_sheet_view_rtl(ws)

        headers = [
            "الدورة المالية", "رقم المشترك", "اسم المشترك",
            "القراءة السابقة", "القراءة الحالية", "الاستهلاك (ك.و)",
            "سعر الكيلو", "رسوم الخدمة", "المتأخرات السابقة",
            "إجمالي المستحق", "رقم الفاتورة", "المبلغ المسدد",
            "رقم سند القبض", "المتبقي المرحل", "حالة الفاتورة"
        ]

        ws.row_dimensions[1].height = 28
        for col_idx, h in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col_idx, value=h)
            cell.fill = self.header_fill
            cell.font = self.header_font
            cell.alignment = self.align_center
            cell.border = self.cell_border

        for row_idx, r in enumerate(sim_result.records, start=2):
            ws.row_dimensions[row_idx].height = 19
            fill = self.zebra_fill if row_idx % 2 == 0 else self.white_fill

            vals = [
                r.cycle_name,
                r.subscriber_number,
                r.full_name,
                r.previous_reading,
                r.current_reading,
                r.consumption_kwh,
                r.kwh_price,
                r.fixed_fee,
                r.carried_arrears,
                r.total_due,
                r.invoice_number or f"INV-{r.cycle_name}-{r.subscriber_number}",
                r.paid_amount,
                r.receipt_number or "-",
                r.remaining_balance,
                r.invoice_status,
            ]

            for c_idx, v in enumerate(vals, start=1):
                cell = ws.cell(row=row_idx, column=c_idx, value=v)
                cell.fill = fill
                cell.border = self.cell_border
                cell.font = Font(name="Segoe UI", size=9)

                if c_idx in (1, 2, 11, 13, 15):
                    cell.alignment = self.align_center
                elif c_idx == 3:
                    cell.alignment = self.align_right
                else:
                    cell.alignment = self.align_left
                    cell.number_format = "#,##0.00"

                if c_idx == 15:
                    if r.invoice_status == "Paid":
                        cell.fill = self.green_fill
                        cell.font = self.green_font
                    elif r.invoice_status == "Unpaid":
                        cell.fill = self.red_fill
                        cell.font = self.red_font

        self._auto_adjust_columns(ws, min_width=13)

    # -------------------------------------------------------------------------
    # Sheet 4: Rollback & Reversal Verification
    # -------------------------------------------------------------------------
    def _create_reversal_audit_sheet(
        self,
        wb: openpyxl.Workbook,
        verification_results: List[VerificationResult],
    ):
        ws = wb.create_sheet(title="4. تدقيق إلغاء السندات")
        self._apply_sheet_view_rtl(ws)

        headers = [
            "م", "معرف السند / الحركة", "نوع الفحص",
            "القيمة المتوقعة (ر.ي)", "القيمة الفعلية (ر.ي)", "فارق الانحراف",
            "النتيجة", "تفاصيل العملية والمطابقة", "الطابع الزمني"
        ]

        ws.row_dimensions[1].height = 28
        for col_idx, h in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col_idx, value=h)
            cell.fill = self.header_fill
            cell.font = self.header_font
            cell.alignment = self.align_center
            cell.border = self.cell_border

        reversals = [v for v in verification_results if v.verification_type == "PAYMENT_REVERSAL"]

        for idx, res in enumerate(reversals, start=1):
            r_num = idx + 1
            ws.row_dimensions[r_num].height = 22
            fill = self.zebra_fill if r_num % 2 == 0 else self.white_fill

            vals = [
                idx,
                res.entity_id,
                res.verification_type,
                res.expected_value,
                res.actual_value,
                res.discrepancy,
                "PASS" if res.passed else "FAIL",
                res.details,
                res.timestamp,
            ]

            for c_idx, v in enumerate(vals, start=1):
                cell = ws.cell(row=r_num, column=c_idx, value=v)
                cell.fill = fill
                cell.border = self.cell_border
                cell.font = Font(name="Segoe UI", size=9)

                if c_idx in (1, 2, 3, 7, 9):
                    cell.alignment = self.align_center
                elif c_idx in (4, 5, 6):
                    cell.alignment = self.align_left
                    cell.number_format = "#,##0.00"
                else:
                    cell.alignment = self.align_right

                if c_idx == 7:
                    cell.fill = self.green_fill if res.passed else self.red_fill
                    cell.font = self.green_font if res.passed else self.red_font

        self._auto_adjust_columns(ws, min_width=14)

    # -------------------------------------------------------------------------
    # Sheet 5: Tariff Escalation Audit
    # -------------------------------------------------------------------------
    def _create_tariff_audit_sheet(
        self,
        wb: openpyxl.Workbook,
        verification_results: List[VerificationResult],
    ):
        ws = wb.create_sheet(title="5. تدقيق تغيير التعرفة")
        self._apply_sheet_view_rtl(ws)

        headers = [
            "م", "معرف الفاتورة", "نوع الفحص",
            "إجمالي الفاتورة المتوقع (ر.ي)", "إجمالي الفاتورة الفعلي (ر.ي)",
            "فارق الانحراف", "النتيجة", "تفاصيل الفحص", "الطابع الزمني"
        ]

        ws.row_dimensions[1].height = 28
        for col_idx, h in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col_idx, value=h)
            cell.fill = self.header_fill
            cell.font = self.header_font
            cell.alignment = self.align_center
            cell.border = self.cell_border

        tariffs = [v for v in verification_results if v.verification_type == "TARIFF_CASCADE"]

        for idx, res in enumerate(tariffs, start=1):
            r_num = idx + 1
            ws.row_dimensions[r_num].height = 22
            fill = self.zebra_fill if r_num % 2 == 0 else self.white_fill

            vals = [
                idx,
                res.entity_id,
                res.verification_type,
                res.expected_value,
                res.actual_value,
                res.discrepancy,
                "PASS" if res.passed else "FAIL",
                res.details,
                res.timestamp,
            ]

            for c_idx, v in enumerate(vals, start=1):
                cell = ws.cell(row=r_num, column=c_idx, value=v)
                cell.fill = fill
                cell.border = self.cell_border
                cell.font = Font(name="Segoe UI", size=9)

                if c_idx in (1, 2, 3, 7, 9):
                    cell.alignment = self.align_center
                elif c_idx in (4, 5, 6):
                    cell.alignment = self.align_left
                    cell.number_format = "#,##0.00"
                else:
                    cell.alignment = self.align_right

                if c_idx == 7:
                    cell.fill = self.green_fill if res.passed else self.red_fill
                    cell.font = self.green_font if res.passed else self.red_font

        self._auto_adjust_columns(ws, min_width=14)

    # -------------------------------------------------------------------------
    # Sheet 6: Anomaly & Discrepancy Log
    # -------------------------------------------------------------------------
    def _create_anomaly_log_sheet(
        self,
        wb: openpyxl.Workbook,
        sim_result: SimulationRunResult,
    ):
        ws = wb.create_sheet(title="6. كشف الشذوذ والتناقضات")
        self._apply_sheet_view_rtl(ws)

        headers = ["م", "نوع الانحراف", "الدورة", "معرف المشترك", "القيمة المتوقعة", "القيمة الفعلية", "الفارق"]
        ws.row_dimensions[1].height = 28
        for col_idx, h in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col_idx, value=h)
            cell.fill = self.header_fill
            cell.font = self.header_font
            cell.alignment = self.align_center
            cell.border = self.cell_border

        if not sim_result.anomalies:
            ws.merge_cells("A2:G3")
            cell = ws["A2"]
            cell.value = "✅ سجل نظيف 100%: لم يتم تسجيل أي انحراف مالي أو تناقض في الحسابات طوال فترة المحاكاة."
            cell.fill = self.green_fill
            cell.font = Font(name="Segoe UI", size=11, bold=True, color="047857")
            cell.alignment = self.align_center
        else:
            for idx, a in enumerate(sim_result.anomalies, start=1):
                r_num = idx + 1
                ws.row_dimensions[r_num].height = 20
                vals = [
                    idx,
                    a.get("type", "UNKNOWN"),
                    a.get("cycle", "-"),
                    a.get("customer_id", "-"),
                    a.get("expected", 0.0),
                    a.get("actual", 0.0),
                    abs(a.get("expected", 0.0) - a.get("actual", 0.0)),
                ]
                for c_idx, v in enumerate(vals, start=1):
                    c = ws.cell(row=r_num, column=c_idx, value=v)
                    c.fill = self.red_fill
                    c.font = self.red_font
                    c.border = self.cell_border
                    c.alignment = self.align_center

        self._auto_adjust_columns(ws, min_width=15)

    # -------------------------------------------------------------------------
    # Master Export Method
    # -------------------------------------------------------------------------
    def export_audit_workbook(
        self,
        sim_result: SimulationRunResult,
        snapshots: List[SubscriberSnapshot],
        verification_results: List[VerificationResult],
        output_path: Optional[Path] = None,
    ) -> Path:
        """Generate complete 6-sheet financial audit workbook and save to disk."""
        target_path = output_path or (
            self.config.reports_dir / f"SmartPower_Financial_Audit_Report_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        )
        target_path.parent.mkdir(parents=True, exist_ok=True)

        wb = openpyxl.Workbook()
        # Remove default sheet
        wb.remove(wb.active)

        logger.info("Building 6-sheet financial audit Excel report at %s...", target_path)

        # 1. Executive Summary
        self._create_executive_summary_sheet(wb, sim_result, verification_results)

        # 2. Subscriber Annual Summary
        self._create_subscriber_summary_sheet(wb, sim_result, snapshots)

        # 3. Monthly Cycle Ledger
        self._create_cycle_ledger_sheet(wb, sim_result)

        # 4. Rollback & Reversal Verification
        self._create_reversal_audit_sheet(wb, verification_results)

        # 5. Tariff Escalation Audit
        self._create_tariff_audit_sheet(wb, verification_results)

        # 6. Anomaly & Discrepancy Log
        self._create_anomaly_log_sheet(wb, sim_result)

        wb.save(str(target_path))
        logger.info("Financial audit report successfully exported to: %s", target_path)
        return target_path
